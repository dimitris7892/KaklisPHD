import numpy as np
import pandas as pd
from scipy.stats import ks_2samp,chisquare,chi2_contingency
import math
#from pyproj import Proj, transform
import pyearth as sp
import matplotlib
from sklearn.cluster import KMeans
from matplotlib.lines import Line2D
from decimal import Decimal
import random
from numpy import inf
#from coordinates.converter import CoordinateConverter, WGS84, L_Est97
import pyodbc
import csv
import pytz
import locale
import time
#locale.setlocale(locale.LC_NUMERIC, "en_DK.UTF-8")
import datetime
import cx_Oracle
from dateutil.rrule import rrule, DAILY, MINUTELY
from sympy.solvers import solve
from sympy import Symbol
from pathlib import Path
import itertools
from sympy import cos, sin , tan , exp , sqrt , E
from openpyxl import load_workbook
import glob, os
from pathlib import Path
#from openpyxl.styles.colors import YELLOW
from openpyxl.styles import Font
from openpyxl.styles.borders import Border, Side
import shutil
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment
#import matplotlib.pyplot as plt
#import seaborn as sns
from openpyxl.drawing.image import Image
from scipy.stats import ttest_ind_from_stats
#pytz.timezone("Eastern European Time")
import geopy
from geopy import  geocoders
from tzwhere import tzwhere
import json

class DBconnections:

    def __init__(self):
        self.connection = None
        self.cursor = None
        self.table = None

    class DBcredentials:

        def __init__(self, user, pwd, server, db):

            self.user = user
            self.pwd = pwd
            self.server = server
            self.db = db

    def connectToDB(self, creds):

        user = creds.user
        pwd = creds.pwd
        server = creds.server
        db = creds.db
        connINSTALLATIONS = pyodbc.connect('DRIVER={ODBC Driver 17 for SQL Server};SERVER='+server+';'
                                            'DATABASE='+db+';'
                                            'UID='+user+';'
                                            'PWD='+pwd+'')

        cursor = connINSTALLATIONS.cursor()

        self.connection =  connINSTALLATIONS
        self.cursor  = cursor
        self.db =  db

        return self.connection

    def createTables(self, company, vessel, type):

        cursor_myserver = self.connection.cursor()

        if type == "tlg":
            if cursor_myserver.tables(table=vessel + '_TLG', tableType='TABLE').fetchone():
                print("TABLE " + vessel + '_TLG' + " ON DATABASE " + str.upper(company) + " EXISTS! ")
            else:
                cursor_myserver.execute(
                    'CREATE TABLE [dbo].[' + vessel + '_TLG' + ']('
                    + '[id] [int] IDENTITY(1,1) NOT NULL,'
                    + '[installationsId] [int] NOT NULL,'
                    + '[vdate] [datetime] NOT NULL,'
                    + '[time_utc] [datetime] NULL,'
                    + '[type] [char](10) NULL,'
                    + '[actl_spd] [decimal](5, 2) NULL,'
                    + '[avg_spd] [decimal](5, 2) NULL,'
                    + '[slip] [decimal](5, 2) NULL,'
                    + '[act_rpm] [decimal](5, 2) NULL,'
                    + '[sail_time] [decimal](5, 2) NULL,'
                    + '[weather_force] [int] NULL,'
                    + '[weather_wind_tv] [char](10) NULL,'
                    + '[sea_condition] [char](10) NULL,'
                    + '[swell_dir] [char](10) NULL,'
                    + '[swell_lvl] [char](10) NULL,'
                    + '[current_dir] [char](10) NULL,'
                    + '[current_speed] [decimal](5, 2) NULL,'
                    + '[draft] [decimal](5, 2) NULL, '
                    + '[lat] [decimal](9, 6) NOT NULL,'
                    + '[lon] [decimal](9, 6) NOT NULL '
                    + 'CONSTRAINT [PK_' + vessel + '_TLG' + '] PRIMARY KEY CLUSTERED '
                    + '('
                    + '     [id] ASC'
                    + ')'
                    + ')'
                )

                # + ') WITH (PAD_INDEX = OFF, STATISTICS_NORECOMPUTE = OFF, IGNORE_DUP_KEY = OFF, ALLOW_ROW_LOCKS = ON, ALLOW_PAGE_LOCKS = ON) ON [PRIMARY]'
                # + ' ) ON [PRIMARY]')

                self.connection.commit()
                print("TABLE " + vessel +'_TLG' " ON DATABASE " + str.upper(company) + " CREATED! ")
        elif type=='raw':
            if cursor_myserver.tables(table=vessel, tableType='TABLE').fetchone():
                print("TABLE " + vessel + " ON DATABASE " + str.upper(company) + " EXISTS! ")
            else:
                cursor_myserver.execute(
                    'CREATE TABLE [dbo].[' + vessel + ']('
                    + '[id] [int] IDENTITY(1,1) NOT NULL,'
                    + '[installationsId] [int] NOT NULL,'
                    + '[lat] [decimal](9, 6) NOT NULL,'
                    + '[lon] [decimal](9, 6) NOT NULL,'
                    + '[t_stamp] [datetime] NULL,'
                    + '[course] [decimal](5, 2) NULL,'
                    + '[stw] [decimal](5, 2) NULL,'
                    + '[s_ovg] [decimal](5, 2) NULL,'
                    + '[me_foc] [decimal](5, 2) NULL,'
                    + '[rpm] [decimal](5, 2) NULL,'
                    + '[power] [int] NULL,'
                    + '[draft] [decimal](5, 2) NULL,'
                    + '[trim] [decimal](5, 2) NULL,'
                    + '[at_port] [int] NULL,'
                    + '[daysFromLastPropClean] [int] NULL,'
                    + '[daysFromLastDryDock] [int] NULL,'
                    + '[sensor_wind_speed] [decimal](5, 2) NULL,'
                    + '[sensor_wind_dir] [decimal](5, 2) NULL,'
                    + '[wind_speed] [decimal](5, 2) NULL,'
                    + '[wind_dir] [decimal](5, 2) NULL,'
                    + '[curr_speed] [decimal](5, 2) NULL,'
                    + '[curr_dir] [decimal](5, 2) NULL,'
                    + '[comb_waves_height] [decimal](5, 2) NULL,'
                    + '[comb_waves_dir] [decimal](5, 2) NULL,'
                    + '[swell_height] [decimal](5, 2) NULL,'
                    + '[swell_dir] [decimal](5, 2) NULL,'
                    + '[wind_waves_height] [decimal](5, 2) NULL,'
                    + '[wind_waves_dir] [decimal](5, 2) NULL, '
                    + '[tlg_id] [int] NULL '
                    + 'CONSTRAINT [PK_' + vessel + '] PRIMARY KEY CLUSTERED '
                    + '('
                    + '     [id] ASC'
                    + ')'
                    + ')'
                )

                # + ') WITH (PAD_INDEX = OFF, STATISTICS_NORECOMPUTE = OFF, IGNORE_DUP_KEY = OFF, ALLOW_ROW_LOCKS = ON, ALLOW_PAGE_LOCKS = ON) ON [PRIMARY]'
                # + ' ) ON [PRIMARY]')

                self.connection.commit()
                print("TABLE " + vessel + " ON DATABASE " + str.upper(company) + " CREATED! ")
        elif type=='processed':
            if cursor_myserver.tables(table=vessel+'_PROCESSED', tableType='TABLE').fetchone():
                print("TABLE " + vessel +'_PROCESSED'+ " ON DATABASE " + str.upper(company) + " EXISTS! ")
            else:
                cursor_myserver.execute(
                    'CREATE TABLE [dbo].[' + vessel +'_PROCESSED'+ ']('
                    + '[id] [int] IDENTITY(1,1) NOT NULL,'
                    + '[installationsId] [int] NOT NULL,'
                    + '[lat] [decimal](9, 6) NOT NULL,'
                    + '[lon] [decimal](9, 6) NOT NULL,'
                    + '[t_stamp] [datetime] NULL,'
                    + '[course] [decimal](5, 2) NULL,'
                    + '[stw] [decimal](5, 2) NULL,'
                    + '[s_ovg] [decimal](5, 2) NULL,'
                    + '[me_foc] [decimal](5, 2) NULL,'
                    + '[rpm] [decimal](5, 2) NULL,'
                    + '[power] [int] NULL,'
                    + '[draft] [decimal](5, 2) NULL,'
                    + '[trim] [decimal](5, 2) NULL,'
                    + '[at_port] [int] NULL,'
                    + '[daysFromLastPropClean] [int] NULL,'
                    + '[daysFromLastDryDock] [int] NULL,'
                    + '[sensor_wind_speed] [decimal](5, 2) NULL,'
                    + '[sensor_wind_dir] [decimal](5, 2) NULL,'
                    + '[wind_speed] [decimal](5, 2) NULL,'
                    + '[wind_dir] [decimal](5, 2) NULL,'
                    + '[curr_speed] [decimal](5, 2) NULL,'
                    + '[curr_dir] [decimal](5, 2) NULL,'
                    + '[comb_waves_height] [decimal](5, 2) NULL,'
                    + '[comb_waves_dir] [decimal](5, 2) NULL,'
                    + '[swell_height] [decimal](5, 2) NULL,'
                    + '[swell_dir] [decimal](5, 2) NULL,'
                    + '[wind_waves_height] [decimal](5, 2) NULL,'
                    + '[wind_waves_dir] [decimal](5, 2) NULL, '
                    + '[tlg_id] [int] NULL '
                    + 'CONSTRAINT [PK_' + vessel +'_PROCESSED'+ '] PRIMARY KEY CLUSTERED '
                    + '('
                    + '     [id] ASC'
                    + ')'
                    + ')'
                )

                # + ') WITH (PAD_INDEX = OFF, STATISTICS_NORECOMPUTE = OFF, IGNORE_DUP_KEY = OFF, ALLOW_ROW_LOCKS = ON, ALLOW_PAGE_LOCKS = ON) ON [PRIMARY]'
                # + ' ) ON [PRIMARY]')

                self.connection.commit()
                print("TABLE " + vessel +'_PROCESSED' " ON DATABASE " + str.upper(company) + " CREATED! ")

    def findLATLONFromRaw(self, company, table, datetime):

        lat , lon = -1, -1

        self.cursor.execute(
            'SELECT lat, lon FROM [dbo].['+table+'] WHERE t_stamp =  '"'" + datetime + "'"'   ')

        for row in self.cursor.fetchall():
            lat, lon = float(row[0]) , float(row[1])

        return lat, lon

    def mapWithTLG(self, company, table, dateTimeRaw):

        tlg_id = -1
        dateRaw = str(dateTimeRaw).split(" ")[0]
        #self.cursor.execute(
            #'SELECT id FROM [dbo].['+table+'_TLG] WHERE  FORMAT (vdate, '"'yyyy-MM-dd'"') =  '"'" + dateRaw + "'"'   ')
        self.cursor.execute('SELECT TOP 1 id '
                            'FROM [dbo].['+table+']'
                            'WHERE [dbo].['+table+'].vdate <  '"'" + dateTimeRaw + "'"
                            'ORDER BY [dbo].['+table+'].vdate DESC')
        rows = self.cursor.fetchall()
        for row in rows:
            tlg_id = int(row[0])

        return tlg_id

    def insertInToTable(self,table,values):

        try:
            self.cursor.execute(
                "insert into [dbo].["+table+"]"
                "(companyName, dateCreated, lastUpdated )"
                + "values ("
                + "'" + str(values.compName) + "', "
                + "'"+ str(values.dateCreated) + "', "
                + "'"+ str(datetime.datetime.now().strftime("%Y%m%d %H:%M:%S")) + "' )")

            self.connection.commit()

        except Exception as e:
            print(str(e))

    def insertInToTableVessel(self, table,compName, values):


        self.cursor.execute(
            'SELECT companyId FROM COMPANIES WHERE companyName =  '"'" + compName + "'"'   ')
        # if cursor_myserver.fetchall().__len__() > 0:
        for row in self.cursor.fetchall():
            companyId = row[0]
            #print(str(imo))

        try:
            self.cursor.execute(
                "insert into [dbo].["+table+"]"
                "( companyId, imo, vesselName, vesselType, dateCreated, dataTypes )"
                + "values ("
                + "'" + str(companyId) + "', "
                + "'" + values.imo + "', "
                + "'"+ values.vslName + "', "
                + "'" + values.vslType + "', "
                + "'"+ str(datetime.datetime.now().strftime("%Y%m%d %H:%M:%S")) +"', "
                + str(values.dataTypes) +
                " )")

            self.connection.commit()

        except Exception as e:
            print(str(e))

    def updateFiledInTableTLG(self, company, vessel, values ,field):

        for i in range(0, len(values)):

            lat, lon = values[i].lat, values[i].lon

            try:
                if lat != -1:
                    tz = tzwhere.tzwhere(forceTZ=True)
                    gmtTimezone = tz.tzNameAt(lat, lon, forceTZ=True)

                    # gmtTimezone = 9.00
                    # gmtTimezone = "+" + str(int(gmtTimezone)) if gmtTimezone > 0 else "-" + str(int(gmtTimezone))
                    timezone = gmtTimezone
                    # timezone = [k for k in list(pytz.all_timezones) if str(k).__contains__(gmtTimezone)][0]
                    local = pytz.timezone(timezone)
                    naive = datetime.datetime.strptime(values[i].vdate, "%Y-%m-%d %H:%M:%S")
                    local_dt = local.localize(naive, is_dst=None)
                    utc_dt = str(local_dt.astimezone(pytz.utc)).split("+")[0]
                else:
                    utc_dt = values[i].vdate
            except Exception as e:
                print(str(e))
                utc_dt = values[i].vdate

            values[i].time_utc = utc_dt

            self.cursor.execute(
                'update  [dbo].[' + vessel + '_TLG] SET '
                                             
                ' '+field+' = '"'" + str(values[i].time_utc) + "'"' '
    
                 'WHERE '
    
                 'lat = '"'" + str(values[i].lat) + "'"' AND '
                 'lon = '"'" + str(values[i].lon) + "'"' AND '
                 'vdate = '"'" + str(values[i].vdate) + "'"'')

        self.connection.commit()
        print("FINISHED  UPDATE ON TABLE " + vessel +'_TLG'  "FOR FILED "+field+" ON DATABASE " + str.upper(company) + "!")

    def insertIntoDataTableforVessel(self, company, vessel ,connInst, type, insertType ,values = None, valuesBulk = None, csv_file_nm = None):

        print("INSERTING INTO " +vessel+ " IN DB "+ company+" . . .")
        connInst.cursor.execute(
            'SELECT vesselId FROM INSTALLATIONS WHERE vesselName =  '"'" + vessel + "'"'   ')
        # if cursor_myserver.fetchall().__len__() > 0:
        for row in connInst.cursor.fetchall():
            installationsId = row[0]


        start_time = time.time()
        if type=='raw' and insertType=='bulk':
            self.cursor.fast_executemany = True
            sql = "insert into [dbo].["+vessel+"] " \
                +"(installationsId, lat, lon, t_stamp , course , stw , s_ovg , me_foc , rpm , power , draft , trim , at_port ,daysFromLastPropClean, "\
                +"daysFromLastDryDock , sensor_wind_speed , sensor_wind_dir , wind_speed ,wind_dir ,curr_speed ,curr_dir ,comb_waves_height ," \
                +"comb_waves_dir , swell_height ,swell_dir, wind_waves_height   , wind_waves_dir, tlg_id ) " \
                +"VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"
            #params = [(f'txt{i:06d}',) for i in range(1000)]
            #qry = "BULK INSERT [dbo].["+vessel+"] FROM '" + csv_file_nm + "' WITH (FORMAT = 'CSV', FIRSTROW = 2)"
            self.cursor.executemany(sql, valuesBulk)

            self.connection.commit()
            print("FINISHED BULK INSERT OR UPDATE ON TABLE " + vessel + " ON DATABASE " + str.upper(company) + "!")

        if type == 'raw' and insertType!='bulk':
            for i in range(0,len(values)):


                tlg_id = self.mapWithTLG(company, vessel, str(values[i].dt))
                self.cursor.execute(
                    'SELECT * FROM [dbo].['+vessel+'] WHERE LAT =  '"'" + str(values[i].OrigLAT) + "'"' AND '
                                                            'LON = '"'" + str(values[i].OrigLON) + "'"' AND '
                                                             't_stamp= '"'" + str(values[i].dt) + "'"'')
                if len(self.cursor.fetchall()) == 0:


                    self.cursor.execute(
                        "insert into [dbo].["+vessel+"](installationsId, lat, lon, t_stamp , course , stw , s_ovg , me_foc , rpm , power , draft , trim , at_port ,daysFromLastPropClean, "
                        + " daysFromLastDryDock , sensor_wind_speed , sensor_wind_dir , wind_speed ,wind_dir ,curr_speed ,curr_dir ,comb_waves_height ,comb_waves_dir , swell_height ,swell_dir,"
                        + "wind_waves_height   , wind_waves_dir, tlg_id ) "

                        + "values ("
                        + str(installationsId) + ", "
                        + str(values[i].OrigLAT) + ", "
                        + str(values[i].OrigLON) + ", "
                        + "'" + str(values[i].dt) + "', "
                        + str(np.nan_to_num(values[i].Vcourse)) + ", "
                        + str(np.nan_to_num(values[i].stw)) + ", "
                        + str(np.nan_to_num(values[i].Sovg)) + " , "
                        + str(np.nan_to_num(values[i].meFoc)) + " , "
                        + str(values[i].rpm) + " , "
                        + str(np.nan_to_num(values[i].power)) + " , "
                        + str(np.nan_to_num(values[i].draft)) + " , "
                        + str(np.nan_to_num(values[i].trim)) + " , "
                        + str(values[i].atPort) + " , "
                        + str(values[i].daysFromLastPropClean) + " , "
                        + str(values[i].daysFromLastDryDock) + " , "
                        + str(np.nan_to_num(values[i].SensorwindSpeed)) + " , "
                        + str(np.nan_to_num(np.round(values[i].relSensorWindDir, 2))) + " , "
                        + str(np.nan_to_num(values[i].windSpeed)) + " , "
                        + str(np.nan_to_num(values[i].relWindDir)) + " , "
                        + str(np.nan_to_num(values[i].currentsSpeed)) + ", "
                        + str(np.nan_to_num(values[i].relCurrentsDir)) + " , "
                        + str(np.nan_to_num(values[i].combSWH)) + " , "
                        + str(np.nan_to_num(values[i].relCombDir)) + " , "
                        + str(np.nan_to_num(values[i].swellSWH)) + " , "
                        + str(np.nan_to_num(values[i].relSwellDir)) + " , "
                        + str(np.nan_to_num(values[i].wavesSWH)) + " , "
                        + str(np.nan_to_num(values[i].relWaveDir)) + " , "
                        + str(tlg_id) + " )")
                else:

                    self.cursor.execute(
                        'update  [dbo].[' + vessel + '] SET '
                        'installationsId = '"'" + str(installationsId) +  "'"', '
                        'course =  '"'" + str(np.nan_to_num(values[i].Vcourse)) +  "'"', '
                        'stw = '"'"  + str(np.nan_to_num(values[i].stw)) +  "'"', '
                        's_ovg = '"'" + str(np.nan_to_num(values[i].Sovg)) +  "'"', '
                        'me_foc = '"'" + str(np.nan_to_num(values[i].meFoc)) +  "'"', '
                        'rpm = '"'" + str(values[i].rpm) +  "'"', '
                        'power = '"'" + str(np.nan_to_num(values[i].power)) +  "'"', '
                        'draft = '"'" + str(np.nan_to_num(values[i].draft)) +  "'"', '
                        'trim = '"'" + str(np.nan_to_num(values[i].trim)) +  "'"', '
                        'at_port = '"'" + str(values[i].atPort) +  "'"','
                        'daysFromLastPropClean = '"'" + str(values[i].daysFromLastPropClean) +  "'"', '
                        'daysFromLastDryDock = '"'" + str(values[i].daysFromLastDryDock) +  "'"', '
                        'sensor_wind_speed = '"'" + str(np.nan_to_num(values[i].SensorwindSpeed)) +  "'"', '
                        'sensor_wind_dir = '"'" + str(np.nan_to_num(np.round(values[i].relSensorWindDir, 2))) +  "'"', '
                        'wind_speed = '"'" + str(np.nan_to_num(values[i].windSpeed)) +  "'"','
                        'wind_dir = '"'" + str(np.nan_to_num(values[i].relWindDir)) +  "'"','
                        'curr_speed = '"'" + str(np.nan_to_num(values[i].currentsSpeed)) +  "'"','
                        'curr_dir = '"'" + str(np.nan_to_num(values[i].relCurrentsDir)) +  "'"','
                        'comb_waves_height = '"'" + str(np.nan_to_num(values[i].combSWH)) +  "'"','
                        'comb_waves_dir = '"'" + str(np.nan_to_num(values[i].relCombDir)) +  "'"','
                        'swell_height = '"'" + str(np.nan_to_num(values[i].swellSWH)) +  "'"','
                        'swell_dir = '"'" + str(np.nan_to_num(values[i].relSwellDir)) +  "'"','
                        'wind_waves_height   = '"'" + str(np.nan_to_num(values[i].wavesSWH)) +  "'"', '
                        'wind_waves_dir = '"'" + str(np.nan_to_num(values[i].relWaveDir)) + "'"', '
                        'tlg_id = '"'" + str(tlg_id) + "'"' '
                                                                                            
                                                                        
                        'WHERE '

                        'LAT = '"'" + str(values[i].OrigLAT) + "'"' AND '
                        'LON = '"'" + str(values[i].OrigLON) + "'"' AND '
                        't_stamp = '"'" + str(values[i].dt) + "'"'')

            self.connection.commit()
            print("FINISHED INSERT OR UPDATE ON TABLE " + vessel + " ON DATABASE " + str.upper(company) + "!")


        elif type=='tlg':
            try:
                for i in range(0, len(values)):
                    self.cursor.execute(
                            "insert into [dbo].[" + vessel + '_TLG' + "](installationsId, vdate, time_utc , type ,actl_spd  , avg_spd , slip , act_rpm , sail_time , weather_force , weather_wind_tv , sea_condition , swell_dir,"
                            + "swell_lvl   , current_dir , current_speed , lat, lon, draft  ) "
                            + "values ("
                            + str(installationsId) + ", "
                            + "'" + str(values[i].vdate) + "', "
                            + "'" + str(values[i].time_utc) + "', "
                            + "'" + str(values[i].tlg_type) + "', "
                            + str(np.nan_to_num(values[i].actl_speed)) + ", "
                            + str(np.nan_to_num(values[i].avg_speed)) + ", "
                            + str(np.nan_to_num(float(values[i].slip))) + ", "
                            + str(np.nan_to_num(values[i].act_rpm)) + " , "
                            + str(np.nan_to_num(values[i].sail_time)) + " , "
                            + str(np.nan_to_num(values[i].weather_force)) + " , "
                            + "'" + str(np.nan_to_num(values[i].weather_wind_tv)) + "', "
                            + "'" + str(np.nan_to_num(values[i].sea_condition)) + "', "
                            + "'" + str(np.nan_to_num(values[i].swell_dir)) + "', "
                            + "'" + str(np.nan_to_num(values[i].swell_lvl)) + "', "
                            + "'" + str(np.nan_to_num(values[i].current_dir)) + "', "
                            + str(np.nan_to_num(values[i].current_speed)) + ", "
                            + str(np.nan_to_num(values[i].lat)) + ", "
                            + str(np.nan_to_num(values[i].lon)) + ", "
                            + str(np.nan_to_num(values[i].draft)) + " )")
                    #print(i)
            except Exception as e:
                print(str(e))
                return

            self.connection.commit()
            print("FINISHED INSERT OR UPDATE ON TABLE " + vessel +'_TLG' " ON DATABASE " + str.upper(company) + "!")

        print("--- %s seconds ---" % (time.time() - start_time))
class Values:

    def __init__(self):
        pass

    def  company(self, compName, dateCreated):

        self.compName = compName
        self.dateCreated = dateCreated
        #self.lastUpd = lastUpd

    def vessel(self, vslName, companyId, vslType, dateCreated, imo, dataTypes):

        self.dateCreated = dateCreated
        self.companyId = companyId
        self.vslName = vslName
        self.vslType = vslType
        self.imo  = imo
        self.dataTypes = dataTypes

    def dataVessel(self):

        self.OrigLAT = -1
        self.OrigLON = -1
        self.dt = -1
        self.Vcourse = -1
        self.stw = -1
        self.Sovg = -1
        self.meFoc = -1
        self.rpm = -1
        self.power = -1
        self.draft = -1
        self.trim = -1
        self.atPort= -1
        self.daysFromLastPropClean= -1
        self.daysFromLastDryDock= -1
        self.SensorwindSpeed= -1
        self.roundrelSensorWindDir= -1
        self.windSpeed= -1
        self.relWindDir = -1
        self.relSensorWindDir = -1
        self.currentsSpeed= -1
        self.relCurrentsDir= -1
        self.combSWH= -1
        self.relCombDir= -1
        self.swellSWH= -1
        self.relSwellDir= -1
        self.wavesSWH= -1
        self.relWaveDir= -1

    def tlgDataVessel(self):

        self.vdate = -1
        self.time_utc = -1
        self.hours = -1
        self.minutes = -1
        self.actl_speed = -1
        self.avg_speed = -1
        self.slip = -1
        self.act_rpm = -1
        self.sail_time = -1
        self.weather_force = -1
        self.weather_wind_tv = -1
        self.sea_condition = -1
        self.swell_dir = -1
        self.swell_lvl = -1
        self.current_dir = -1
        self.current_speed = -1
        self.bl_flag = -1
        self.tlg_type = -1
        self.port_name = -1
        self.lat = -1
        self.lon = -1
        self.draft = -1


def fillDataTlgVesselClass(company, vessel):

    tlgs = pd.read_csv('./data/' + company + '/' + vessel + '/TELEGRAMS/'+vessel+'.csv', sep=';').values
    tlgDataVesselClassList = []

    for i in range(0, len(tlgs)):

        values = Values()
        values.tlgDataVessel()

        values.vdate = tlgs[i][0]
        values.actl_speed = tlgs[i][12]
        values.avg_speed = tlgs[i][12]
        #values.slip = None
        values.act_rpm = tlgs[i][9]
        values.sail_time = tlgs[i][7]
        values.weather_force = tlgs[i][11]
        values.weather_wind_tv = tlgs[i][10]
        #values.sea_condition = None
        values.swell_dir = tlgs[i][21]
        values.swell_lvl = tlgs[i][20]
        values.lat = tlgs[i][5] + tlgs[i][4] / 60
        values.lon = tlgs[i][3] + tlgs[i][6] / 60
        #values.current_dir = None
        #values.current_speed = None
        values.port_name = tlgs[i][22]
        values.bl_flag = tlgs[i][2]
        values.tlg_type = tlgs[i][1]
        values.draft = tlgs[i][8]

        #lat, lon = connData.findLATLON(company, vessel.split(" ")[0]+'_'+vessel.split(" ")[1], values.vdate)
        lat, lon = values.lat, values.lon

        try:
            if lat != -1:
                tz = tzwhere.tzwhere(forceTZ=True)
                gmtTimezone = tz.tzNameAt(lat, lon, forceTZ=True)

                # gmtTimezone = 9.00
                # gmtTimezone = "+" + str(int(gmtTimezone)) if gmtTimezone > 0 else "-" + str(int(gmtTimezone))
                timezone = gmtTimezone
                # timezone = [k for k in list(pytz.all_timezones) if str(k).__contains__(gmtTimezone)][0]
                local = pytz.timezone(timezone)
                naive = datetime.datetime.strptime(values.vdate, "%Y-%m-%d %H:%M:%S")
                local_dt = local.localize(naive, is_dst=None)
                utc_dt = str(local_dt.astimezone(pytz.utc)).split("+")[0]
            else:
                utc_dt = values.vdate
        except:
            utc_dt = values.vdate

        values.time_utc = utc_dt
        #values.time_utc = None

        #print(i)
        '''values.swell_lvl = 'n/a' if str(values.swell_lvl) == 'nan' else str(values.swell_lvl)
        try:
            values.slip = str(values.slip).split(',')[0] + "." + str(values.slip).split(',')[1] if str(values.slip).split(
                ',').__len__() > 0 and str(values.slip) != 'nan' else values.slip
        except Exception as e:
            x=0'''
            #print(str(e))


        tlgDataVesselClassList.append(values)

    return tlgDataVesselClassList

def fillDataVesselClass(company, vessel, connInst, connData, dataCleaned = None):

    tableNameVslRaw =  vessel.split(" ")[0]+"_"+vessel.split(" ")[1]
    tableNameVsl_TLG = vessel.split(" ")[0]+"_"+vessel.split(" ")[1]+"_TLG"
    tableNameVsl = tableNameVslRaw +"_PROCESSED" if len(dataCleaned) > 0 else tableNameVslRaw
    ##find InstallaitonId
    connInst.cursor.execute(
        'SELECT vesselId FROM INSTALLATIONS WHERE vesselName =  '"'" +tableNameVslRaw+ "'"'   ')
    # if cursor_myserver.fetchall().__len__() > 0:
    for row in connInst.cursor.fetchall():
        installationsId = row[0]
    #########################################################

    data = pd.read_csv('./data/'+company+'/'+vessel+'/mappedDataNew.csv').values if len(dataCleaned) == 0 else dataCleaned
    tlgs = pd.read_csv('./data/' + company + '/' + vessel + '/TELEGRAMS/'+vessel+'.csv', sep=';').values
    dataVesselClassList = []
    dataVesselClassListBulk = []
    for i in range(0, len(data)):

        values = Values()
        values.dataVessel()
        values.dt = data[i][0]
        ##find if at port
        date = str(values.dt).split(" ")[0]

        telegramRows = np.array([row for row in tlgs if str(row[0]).split(" ")[0] == date])
        if len(telegramRows) > 0:
            tlgsDates = [datetime.datetime.strptime( k , '%Y-%m-%d %H:%M:%S') for k in telegramRows[:,0]]
            closestDate = min(tlgsDates, key=lambda x: abs(x - datetime.datetime.strptime(values.dt, '%Y-%m-%d %H:%M:%S')))
            tlgRow = telegramRows[telegramRows[:,0]==str(closestDate)][0]
            tlgType = tlgRow[1]
            values.atPort = 1 if tlgType == 'A' or tlgType =='D' else 0
        else:
            values.atPort = 1


        values.Vcourse = data[i][1]
        values.blFlag = data[i][2]
        values.draft = data[i][8]
        values.relSensorWindDir = data[i][10]
        values.SensorwindSpeed = data[i][11]
        values.stw = data[i][12]
        values.meFoc = data[i][15]
        values.trim = data[i][17]
        values.swellSWH = data[i][20]
        values.relSwellDir = data[i][21]
        values.combSWH = data[i][22]
        values.relCombDir = data[i][23]
        values.wavesSWH = data[i][24]
        values.relWaveDir = data[i][25]
        values.OrigLAT = data[i][26]
        values.OrigLON = data[i][27]
        values.windSpeed = data[i][28]
        values.relWindDir = data[i][27]
        #values.currentsSpeed = data[i][29]
        #values.relCurrentsDir = data[i][30]
        #values.Sovg = data[i][31]
        tlg_id = connData.mapWithTLG(company, tableNameVsl_TLG, str(values.dt))

        dataVesselClassListBulk.append((
            installationsId,
            values.OrigLAT,
            values.OrigLON,
            values.dt,
            values.Vcourse ,
            values.stw,
            values.Sovg,
            values.meFoc,
            values.rpm,
            values.power,
            values.draft,
            values.trim,
            values.atPort,
            values.daysFromLastPropClean,
            values.daysFromLastDryDock,
            values.SensorwindSpeed,
            values.relSensorWindDir,
            values.windSpeed,
            values.relWindDir,
            values.currentsSpeed,
            values.relCurrentsDir,
            values.combSWH,
            values.relCombDir,
            values.swellSWH,
            values.relSwellDir,
            values.wavesSWH,
            values.relWaveDir,
            tlg_id
        ))  # append data


        dataVesselClassList.append(values)

    return dataVesselClassList , dataVesselClassListBulk


def main():

    connInst = DBconnections()
    instanceInfoInstallationsDB = connInst.DBcredentials("sa", "Man1f0ld", 'localhost, 1433', 'Installations')
    connInst.connectToDB(instanceInfoInstallationsDB)

    connData = DBconnections()
    instanceDataDB = connData.DBcredentials("sa", "Man1f0ld", 'localhost, 1433', 'DANAOS')
    cnxn = connData.connectToDB(instanceDataDB)

    connData.createTables("DANAOS", "EXPRESS_ATHENS", 'tlg')
    connData.createTables("DANAOS","EXPRESS_ATHENS",'raw')
    connData.createTables("DANAOS", "EXPRESS_ATHENS", 'processed')
    #return
    values = Values()
    #values.company("DANAOS",datetime.datetime.now().strftime("%Y%m%d %H:%M:%S"))
    #values.vessel("EXPRESS ATHENS", "" ,"Container", datetime.datetime.now().strftime("%Y%m%d %H:%M:%S"), "9484948", 2)

    #values.dataVessel()
    #conn.insertInToTable("COMPANIES",values)

    #conn.insertInToTableVessel("INSTALLATIONS", "DANAOS" ,values)

    dataVesselClassList, dataVesselClassListBulk = fillDataVesselClass("DANAOS","EXPRESS ATHENS", connInst, connData)
    #dataVesselClassList = fillDataTlgVesselClass("DANAOS", "EXPRESS ATHENS")


    #connData.insertIntoDataTableforVessel("DANAOS","EXPRESS_ATHENS",dataVesselClassList, connInst, "tlg")

    #connData.updateFiledInTableTLG("DANAOS", "EXPRESS_ATHENS", dataVesselClassList, 'time_utc')


    connData.insertIntoDataTableforVessel("DANAOS", "EXPRESS_ATHENS", dataVesselClassList, connInst, "raw","bulk", dataVesselClassListBulk)

if __name__ == "__main__":
    main()